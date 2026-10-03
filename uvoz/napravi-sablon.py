# -*- coding: utf-8 -*-
"""Pravi Excel sablon za masovni uvoz (aktuelni + istorijski projekti) u GradnjaOS.
Kolone 1:1 prate supabase/schema.sql; padajuci meniji za sve sifrarnike; po jedan primer
reda na svakom listu (zuto = popunjava korisnik). Pokretanje: python uvoz/napravi-sablon.py
"""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'GradnjaOS-uvoz-sablon.xlsx')
FONT = 'Arial'
H_FILL = PatternFill('solid', fgColor='1F4E79'); H_FONT = Font(name=FONT, bold=True, color='FFFFFF', size=10)
IN_FILL = PatternFill('solid', fgColor='FFF9C4')       # zuto: popunjava korisnik
REQ_FILL = PatternFill('solid', fgColor='FFE0B2')      # narandzasto: obavezno
EX_FONT = Font(name=FONT, italic=True, color='666666', size=10)
THIN = Side(style='thin', color='D0D0D0'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ---------- sifrarnici (isti kao u index.html / schema.sql) ----------
S = {
    'modul': ['izvodjenje', 'projektovanje'],
    'tip': ['visokogradnja', 'niskogradnja'],
    'nivo': ['IDR', 'IDP/PGD', 'PZI'],
    'status_gr': ['planirano', 'u toku', 'kasni', 'zavrseno'],
    'faza': ['Priprema terena', 'Zemljani radovi', 'Temelji', 'Konstrukcija', 'Grubi radovi', 'Instalacije', 'Završni radovi',
             'Tehnički prijem', 'Odvodnjavanje', 'Donji stroj', 'Gornji stroj', 'Signalizacija', 'Predato',
             'Idejno rešenje', 'IDR', 'IDP/PGD', 'PZI', 'Tehnička kontrola'],
    'adm_st': ['nema', 'priprema', 'ima'],
    'tip_klijenta': ['Investitor', 'Javni sektor', 'Fizičko lice'],
    'pozicija': ['Rukovodilac gradilišta', 'Građevinski inženjer', 'Arhitekta', 'Inženjer geodezije', 'Inženjer instalacija',
                 'Građevinski tehničar', 'Nadzorni organ', 'Majstor', 'Pomoćni radnik'],
    'status_zap': ['Na terenu', 'Kancelarija', 'Odsutan'],
    'da_ne': ['DA', 'NE'],
    'jm': ['m³', 'm²', 'm', 'kom', 'kg', 't', 'h', 'paušal'],
    'prio': ['high', 'mid', 'low'],
    'kolona': ['todo', 'inprogress', 'hold', 'done'],
    'kat': ['Materijal', 'Radna snaga', 'Mehanizacija', 'Podizvođači', 'Ostalo'],
    'status_sit': ['ceka', 'placeno', 'kasni'],
    'status_pod': ['aktivan', 'zavrsen'],
    'tip_res': ['vozilo', 'alat', 'racunar', 'licenca', 'polisa'],
}

# ---------- listovi: (naziv, [ (kolona, obavezno, sifrarnik|None, tip, opis) ], primer) ----------
D = 'datum GGGG-MM-DD (ili Excel datum)'
SHEETS = [
    ('Klijenti', [
        ('id', False, None, 'tekst', 'Šifra (npr. KL-01). Prazno = generiše se.'),
        ('naziv', True, None, 'tekst', 'Naziv investitora / klijenta'),
        ('tip', False, 'tip_klijenta', 'lista', ''),
        ('osoba', False, None, 'tekst', 'Kontakt osoba'),
        ('tel', False, None, 'tekst', ''),
        ('mail', False, None, 'tekst', ''),
    ], ['KL-01', 'Delta Real Estate d.o.o.', 'Investitor', 'Marko Marković', '011 123 4567', 'marko@delta.rs']),
    ('Zaposleni', [
        ('id', False, None, 'tekst', 'Šifra (npr. Z-01). Prazno = generiše se. Koristi se u drugim listovima.'),
        ('ime', True, None, 'tekst', 'Ime i prezime'),
        ('poz', False, 'pozicija', 'lista', 'Pozicija'),
        ('status', False, 'status_zap', 'lista', ''),
        ('tel', False, None, 'tekst', ''),
        ('opis', False, None, 'tekst', 'Opis posla'),
    ], ['Z-01', 'Petar Kovačević', 'Rukovodilac gradilišta', 'Na terenu', '064 111 2222', 'Vodi gradilište, koordinira podizvođače']),
    ('Gradilišta', [
        ('id', False, None, 'tekst', 'Šifra projekta (npr. G-01). Prazno = generiše se.'),
        ('naziv', True, None, 'tekst', 'Naziv projekta'),
        ('modul', True, 'modul', 'lista', 'izvodjenje | projektovanje'),
        ('tip', False, 'tip', 'lista', 'Samo za izvođenje: visokogradnja | niskogradnja'),
        ('nivo', False, 'nivo', 'lista', 'Samo za projektovanje: IDR | IDP/PGD | PZI'),
        ('lok', False, None, 'tekst', 'Lokacija'),
        ('klijent_id', False, None, 'tekst', 'id iz lista Klijenti'),
        ('rukovodilac_id', True, None, 'tekst', 'id iz lista Zaposleni — ko vidi projekat kao rukovodilac'),
        ('pocetak', True, None, 'datum', D),
        ('rok', True, None, 'datum', D),
        ('napredak', False, None, 'broj', '0–100 (%). Istorijski/završen projekat = 100'),
        ('status', True, 'status_gr', 'lista', 'Istorijski projekat = zavrseno'),
        ('faza', False, 'faza', 'lista', 'Trenutna faza'),
        ('povrsina', False, None, 'broj', 'm²'),
        ('nadzor', False, None, 'tekst', 'Kontakt nadzora (ime, telefon)'),
        ('budzet', False, None, 'broj', 'Ugovorena cena (€) — vidi samo direktor'),
        ('troskovi', False, None, 'broj', 'Planirani troškovi (€)'),
        ('potroseno', False, None, 'broj', 'Utrošeno (€) — koristi se samo ako nema stavki u listu Troškovi'),
        ('naplaceno', False, None, 'broj', 'Naplaćeno (€) — koristi se samo ako nema stavki u listu Situacije'),
        ('ugovor_status', False, 'adm_st', 'lista', 'Ugovor / radni nalog: nema | priprema | ima'),
        ('ugovor_vazi_do', False, None, 'datum', D),
        ('prijava_status', False, 'adm_st', 'lista', 'Prijava radova'),
        ('prijava_vazi_do', False, None, 'datum', D),
        ('polisa_status', False, 'adm_st', 'lista', 'Polisa osiguranja'),
        ('polisa_vazi_do', False, None, 'datum', D),
        ('podugovori_status', False, 'adm_st', 'lista', 'Ugovori podizvođača / saradnika'),
    ], ['G-01', 'Stambeni kompleks „Liman Park"', 'izvodjenje', 'visokogradnja', '', 'Novi Sad, Liman', 'KL-01', 'Z-01',
        '2025-03-01', '2026-12-31', 62, 'u toku', 'Konstrukcija', 4200, 'Ing. M. Petrović, 063 555 111', 2450000, 2050000, '', '',
        'ima', '2026-12-31', 'ima', '', 'ima', '2026-06-30', 'priprema']),
    ('Tim', [
        ('zaposleni_id', True, None, 'tekst', 'id iz lista Zaposleni'),
        ('gradiliste_id', True, None, 'tekst', 'id iz lista Gradilišta'),
        ('aktivan', True, 'da_ne', 'lista', 'DA = trenutno radi; NE = ranije radio (istorija)'),
    ], ['Z-01', 'G-01', 'DA']),
    ('Predmer', [
        ('id', False, None, 'tekst', 'Prazno = generiše se'),
        ('gradiliste_id', True, None, 'tekst', 'id iz lista Gradilišta'),
        ('poz', True, None, 'tekst', 'Opis pozicije / usluge'),
        ('jm', False, 'jm', 'lista', ''),
        ('kol', False, None, 'broj', 'Ugovorena količina'),
        ('cena', False, None, 'broj', 'Jedinična cena (€)'),
        ('izv', False, None, 'broj', 'Izvedena količina do sada'),
        ('zaduzen_id', False, None, 'tekst', 'id zaposlenog ili podizvođača'),
    ], ['', 'G-01', 'Iskop temelja', 'm³', 180, 22, 180, 'Z-01']),
    ('Zadaci', [
        ('id', False, None, 'tekst', 'Prazno = generiše se'),
        ('gradiliste_id', True, None, 'tekst', ''),
        ('naziv', True, None, 'tekst', ''),
        ('zaduzen_id', False, None, 'tekst', 'id iz lista Zaposleni'),
        ('prio', False, 'prio', 'lista', 'high | mid | low'),
        ('kolona', False, 'kolona', 'lista', 'todo | inprogress | hold | done'),
        ('rok', True, None, 'datum', D),
    ], ['', 'G-01', 'Armiranje ploče 3. sprata', 'Z-01', 'high', 'inprogress', '2026-10-20']),
    ('Dnevnik', [
        ('id', False, None, 'tekst', 'Prazno = generiše se'),
        ('gradiliste_id', True, None, 'tekst', ''),
        ('datum', True, None, 'datum', D),
        ('autor_id', True, None, 'tekst', 'id iz lista Zaposleni'),
        ('tekst', True, None, 'tekst', 'Unos u dnevnik radova'),
    ], ['', 'G-01', '2026-10-03', 'Z-01', 'Betoniranje ploče 3. sprata, 12 radnika, vreme sunčano.']),
    ('Troškovi', [
        ('id', False, None, 'tekst', 'Prazno = generiše se'),
        ('gradiliste_id', True, None, 'tekst', ''),
        ('datum', True, None, 'datum', D),
        ('opis', True, None, 'tekst', ''),
        ('kat', False, 'kat', 'lista', ''),
        ('iznos', True, None, 'broj', '€'),
        ('dobavljac', False, None, 'tekst', ''),
        ('br_fakture', False, None, 'tekst', ''),
    ], ['', 'G-01', '2026-09-28', 'Beton MB30, 48 m³', 'Materijal', 5760, 'Beton d.o.o.', '2026-117']),
    ('Situacije', [
        ('id', False, None, 'tekst', 'Prazno = generiše se'),
        ('gradiliste_id', True, None, 'tekst', ''),
        ('br', True, None, 'tekst', 'Broj situacije, npr. PS-03/26'),
        ('opis', False, None, 'tekst', ''),
        ('iznos', True, None, 'broj', '€'),
        ('izdato', True, None, 'datum', D),
        ('valuta', True, None, 'datum', 'Rok plaćanja'),
        ('status', True, 'status_sit', 'lista', 'ceka | placeno | kasni'),
    ], ['', 'G-01', 'PS-03/26', 'Grubi radovi — 3. sprat', 185000, '2026-09-30', '2026-10-30', 'ceka']),
    ('Podizvođači', [
        ('id', False, None, 'tekst', 'Šifra (npr. P-01). Prazno = generiše se.'),
        ('naziv', True, None, 'tekst', 'Podizvođač / spoljni saradnik'),
        ('delatnost', False, None, 'tekst', ''),
        ('osoba', False, None, 'tekst', ''),
        ('tel', False, None, 'tekst', ''),
        ('cena', False, None, 'broj', 'Ugovorena vrednost (€)'),
        ('status', False, 'status_pod', 'lista', 'aktivan | zavrsen'),
        ('gradilista_ids', False, None, 'tekst', 'Više id-jeva razdvojenih zarezom: G-01, G-02'),
    ], ['P-01', 'Elektro Tim d.o.o.', 'Elektroinstalacije', 'Ivan Ilić', '065 333 4444', 120000, 'aktivan', 'G-01']),
    ('Resursi', [
        ('id', False, None, 'tekst', 'Prazno = generiše se'),
        ('tip', True, 'tip_res', 'lista', 'vozilo | alat | racunar | licenca | polisa'),
        ('naziv', True, None, 'tekst', ''),
        ('oznaka', False, None, 'tekst', 'Registracija / serijski broj / broj polise'),
        ('gradiliste_id', False, None, 'tekst', 'Prazno = baza / kancelarija'),
        ('zaduzen_id', False, None, 'tekst', 'id iz lista Zaposleni'),
        ('istice', False, None, 'datum', 'Istek registracije / licence / polise (prazno = ne ističe)'),
        ('napomena', False, None, 'tekst', ''),
    ], ['', 'vozilo', 'Kombi VW Crafter', 'NS-123-AB', 'G-01', 'Z-01', '2027-02-15', 'Registracija']),
    ('Magacin', [
        ('id', False, None, 'tekst', 'Prazno = generiše se'),
        ('naziv', True, None, 'tekst', 'Artikal'),
        ('jm', False, 'jm', 'lista', ''),
        ('stanje', False, None, 'broj', 'Trenutno stanje'),
    ], ['', 'Cement CEM II 42,5', 'kg', 2400]),
]

UPUTSTVO = [
    ('GradnjaOS — šablon za masovni uvoz projekata', 'naslov'),
    ('Jedan fajl puni celu bazu: aktuelne i istorijske projekte, klijente, zaposlene, timove, predmere, zadatke, dnevnik, troškove, situacije, podizvođače, resurse i magacin.', ''),
    ('', ''),
    ('Kako se popunjava', 'pod'),
    ('1. Popunjavaj ŽUTA polja. NARANDŽASTE kolone su obavezne. Red 2 na svakom listu je PRIMER — obriši ga ili prepiši.', ''),
    ('2. Šifre (id) povezuju listove: Gradilišta.klijent_id → Klijenti.id, Gradilišta.rukovodilac_id → Zaposleni.id, Tim/Predmer/Zadaci/… .gradiliste_id → Gradilišta.id. Ako ostaviš id prazan, uvoz ga generiše (G-01, Z-01…) — ali onda ga ne možeš koristiti u drugim listovima, pa je bolje da ga sam upišeš.', ''),
    ('3. Datumi: GGGG-MM-DD (2026-10-04) ili običan Excel datum. Iznosi u evrima bez simbola; decimala zarez ili tačka.', ''),
    ('4. Padajući meniji (modul, status, faza, kategorija…) — upiši tačno ponuđenu vrednost; uvoz odbija nepoznate.', ''),
    ('5. Istorijski (završeni) projekti: status = zavrseno, napredak = 100, pocetak/rok u prošlosti. Tim za njih upiši u listu Tim sa aktivan = NE (ulazi u istoriju zaposlenog).', ''),
    ('6. Finansije (budzet, troskovi, cena, iznosi) vidi samo direktor — rukovodioci ih ne dobijaju ni preko API-ja.', ''),
    ('7. Modul projektovanje: popuni nivo (IDR / IDP/PGD / PZI); tip ostavi prazan. Modul izvodjenje: popuni tip; nivo prazan.', ''),
    ('8. Troškovi i Situacije po stavkama su izvor istine za „utrošeno" i „naplaćeno"; kolone potroseno/naplaceno na Gradilištima koristi samo ako nemaš stavke.', ''),
    ('', ''),
    ('Šta se dešava pri uvozu', 'pod'),
    ('python uvoz/uvezi.py <fajl.xlsx>  → proverava šifre, veze i padajuće vrednosti, pa pravi uvoz.sql (INSERT … ON CONFLICT DO UPDATE — ponovni uvoz istog fajla samo osvežava, ne duplira).', ''),
    ('SQL se pokreće na bazi (Supabase SQL editor ili kroz Claude) — jednom, u transakciji. Greške u fajlu se prijavljuju PRE bilo kakvog upisa.', ''),
    ('', ''),
    ('Legenda boja', 'pod'),
    ('ŽUTO = polje koje popunjavaš   ·   NARANDŽASTO = obavezna kolona   ·   SIVO kurziv = primer   ·   zaglavlje = naziv kolone u bazi', ''),
]


def main():
    wb = Workbook()
    ws = wb.active; ws.title = 'Uputstvo'
    ws.column_dimensions['A'].width = 140
    for i, (txt, st) in enumerate(UPUTSTVO, 1):
        c = ws.cell(row=i, column=1, value=txt)
        c.font = Font(name=FONT, size=16 if st == 'naslov' else 12 if st == 'pod' else 10, bold=st in ('naslov', 'pod'),
                      color='1F4E79' if st else '000000')
        c.alignment = Alignment(wrap_text=True, vertical='top')
    r = len(UPUTSTVO) + 2
    ws.cell(row=r, column=1, value='Primer boja:').font = Font(name=FONT, bold=True, size=10)
    ws.cell(row=r + 1, column=1, value='ovako izgleda polje koje popunjavaš').fill = IN_FILL
    ws.cell(row=r + 2, column=1, value='ovako izgleda obavezno polje').fill = REQ_FILL

    # Sifarnici
    sf = wb.create_sheet('Šifarnici')
    col = 1
    rng = {}
    for key, vals in S.items():
        sf.cell(row=1, column=col, value=key).font = H_FONT
        sf.cell(row=1, column=col).fill = H_FILL
        for j, v in enumerate(vals, 2):
            sf.cell(row=j, column=col, value=v).font = Font(name=FONT, size=10)
        L = get_column_letter(col)
        rng[key] = f"'Šifarnici'!${L}$2:${L}${len(vals) + 1}"
        sf.column_dimensions[L].width = max(14, max(len(v) for v in vals) + 2)
        col += 1

    for name, cols, primer in SHEETS:
        s = wb.create_sheet(name)
        s.freeze_panes = 'A2'
        for j, (c, req, lst, tip, opis) in enumerate(cols, 1):
            h = s.cell(row=1, column=j, value=c)
            h.font = H_FONT; h.fill = H_FILL; h.alignment = Alignment(vertical='center', wrap_text=True); h.border = BORDER
            if opis or req:
                h.comment = Comment(('OBAVEZNO. ' if req else '') + (opis or ''), 'GradnjaOS')
            L = get_column_letter(j)
            s.column_dimensions[L].width = max(12, min(44, len(c) + 6 if tip != 'tekst' else 24))
            # primer (red 2) + ulazna polja (redovi 3..500)
            ex = s.cell(row=2, column=j, value=(primer[j - 1] if j - 1 < len(primer) else ''))
            ex.font = EX_FONT; ex.border = BORDER
            if tip == 'datum': ex.number_format = 'yyyy-mm-dd'
            for rr in range(3, 501):
                cell = s.cell(row=rr, column=j)
                cell.fill = REQ_FILL if req else IN_FILL
                cell.font = Font(name=FONT, size=10); cell.border = BORDER
                if tip == 'datum': cell.number_format = 'yyyy-mm-dd'
                elif tip == 'broj': cell.number_format = '#,##0.##'
            if lst:
                dv = DataValidation(type='list', formula1=rng[lst], allow_blank=True, showErrorMessage=True,
                                    errorTitle='Nepoznata vrednost', error='Izaberi vrednost iz padajućeg menija (vidi list Šifarnici).')
                s.add_data_validation(dv); dv.add(f'{L}2:{L}500')
            elif tip == 'datum':
                dv = DataValidation(type='date', operator='between', formula1='DATE(1990,1,1)', formula2='DATE(2100,12,31)',
                                    allow_blank=True, showErrorMessage=True, errorTitle='Datum', error='Unesi datum (GGGG-MM-DD).')
                s.add_data_validation(dv); dv.add(f'{L}3:{L}500')
            elif tip == 'broj':
                dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                                    showErrorMessage=True, errorTitle='Broj', error='Unesi broj ≥ 0.')
                s.add_data_validation(dv); dv.add(f'{L}3:{L}500')
        s.row_dimensions[1].height = 30
        s.auto_filter.ref = f'A1:{get_column_letter(len(cols))}500'
    wb.save(OUT)
    print('sablon:', OUT)


if __name__ == '__main__':
    main()
